# New in floodstate-eo, 2026-10-01 (maintainer: "a version of the document with the tables and figures inserted"). STATUS: ACTIVE.
"""p101_assemble_doc -- the manuscript with every figure and table inserted at its first mention, as DOCX (for proofreading in a
word processor) and as Markdown with image links.

Input: publication/manuscript_uk.md (--lang uk, the Ukrainian proofreading copy) or publication/manuscript.md (--lang en); the
figures of publication/figures/<FigNN>_*.png with their captions from publication/captions.md; the tables of
publication/tables/T*.csv with their captions from publication/tables/manifest.json. A figure or table is inserted once, after the
paragraph that first mentions it (panel letters ignored); tables longer than MAX_ROWS rows or wider than MAX_COLS columns are cut
with a note that names the CSV. Images are downscaled copies (JPEG, <= IMG_W px wide) so that the DOCX stays small.
Outputs: publication/assembled/manuscript_<lang>_assembled.{docx,md}, publication/assembled/img/<FigNN>.jpg (the .docx and img/ are
git-ignored; the .md is tracked). Nothing here changes a number: the text is the filled manuscript verbatim.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

HERE = Path(__file__).resolve().parent
CS = HERE.parents[1]; PUB = CS / "publication"; FIGS = PUB / "figures"; TABS = PUB / "tables"; OUT = PUB / "assembled"
MAX_ROWS, MAX_COLS, IMG_W = 25, 10, 1600
FIG_RE = re.compile(r"\b(FigS?\d{2})[a-z]?\b"); TAB_RE = re.compile(r"\bT\d{2}[a-z]{0,2}\b")
NOTE = {"uk": ("Рисунок", "Таблиця", "показано {r} з {R} рядків і {c} з {C} колонок; повна таблиця: publication/tables/{t}.csv", "джерело підпису: captions.md (англійською)"),
        "en": ("Figure", "Table", "{r} of {R} rows and {c} of {C} columns shown; full table: publication/tables/{t}.csv", "")}


def captions():
    """FigID -> caption text, from captions.md (several supplementary captions share one line)."""
    t = (PUB / "captions.md").read_text(encoding="utf-8")
    hits = list(re.finditer(r"\*\*(FigS?\d{2})\b", t)); out = {}
    for i, m in enumerate(hits):
        end = hits[i + 1].start() if i + 1 < len(hits) else len(t)
        cap = re.sub(r"\s+", " ", t[m.start():end]).strip()
        cap = re.sub(r"^\*\*(FigS?\d{2})\*\*\s*", r"**\1.** ", cap)                 # '**FigS01** text' -> '**FigS01.** text'
        out.setdefault(m.group(1), cap)
    return out


def fig_file(fid):
    c = sorted(FIGS.glob(f"{fid}_*.png"))
    return c[0] if c else None


def small_image(src, fid):
    from PIL import Image
    OUT.joinpath("img").mkdir(parents=True, exist_ok=True); dst = OUT / "img" / f"{fid}.jpg"
    im = Image.open(src).convert("RGB")
    if im.width > IMG_W:
        im = im.resize((IMG_W, round(im.height * IMG_W / im.width)), Image.LANCZOS)
    im.save(dst, "JPEG", quality=82, optimize=True); return dst


def table_block(tid, lang):
    import pandas as pd
    df = pd.read_csv(TABS / f"{tid}.csv", dtype=str, keep_default_na=False)
    R, C = df.shape; cut = df.iloc[:MAX_ROWS, :MAX_COLS]
    note = NOTE[lang][2].format(r=min(R, MAX_ROWS), R=R, c=min(C, MAX_COLS), C=C, t=tid) if (R > MAX_ROWS or C > MAX_COLS) else ""
    return cut, note


def md_table(df):
    esc = lambda v: str(v).replace("|", "\\|").replace("\n", " ")
    lines = ["| " + " | ".join(esc(c) for c in df.columns) + " |", "|" + "---|" * len(df.columns)]
    lines += ["| " + " | ".join(esc(v)[:60] for v in row) + " |" for row in df.itertuples(index=False)]
    return "\n".join(lines)


def runs(par, text):
    """Markdown inline **bold** / *italic* / `code` into python-docx runs (good enough for proofreading)."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            par.add_run(tok[1:-1]).italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
        else:
            par.add_run(tok)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--lang", choices=["uk", "en"], default="uk"); a = ap.parse_args()
    import docx
    from docx.shared import Cm, Pt
    src = PUB / ("manuscript_uk.md" if a.lang == "uk" else "manuscript.md"); text = src.read_text(encoding="utf-8")
    caps = captions(); tman = json.loads((TABS / "manifest.json").read_text(encoding="utf-8"))["tables"]
    fig_word, tab_word, _, cap_note = NOTE[a.lang]
    OUT.mkdir(parents=True, exist_ok=True)
    doc = docx.Document(); st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(2.0)
    md = []; seen_f, seen_t = [], []; n_fig = n_tab = 0
    paragraphs = [p for p in re.split(r"\n\s*\n", text)]
    for para in paragraphs:
        p = para.strip("\n")
        if not p.strip():
            continue
        md.append(p)
        if p.startswith("#"):
            level = len(p) - len(p.lstrip("#")); doc.add_heading(p.lstrip("#").strip(), level=min(level - 1, 3) if level > 1 else 0)
        elif re.match(r"^\s*[-*] ", p):
            for item in re.split(r"\n(?=\s*[-*] )", p):
                runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", item).replace("\n", " "))
        elif re.match(r"^\s*\d+\. ", p):
            for item in re.split(r"\n(?=\s*\d+\. )", p):
                runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\. ", "", item).replace("\n", " "))
        elif p.startswith(">"):
            runs(doc.add_paragraph(style="Intense Quote"), re.sub(r"^>\s?", "", p, flags=re.M).replace("\n", " "))
        else:
            runs(doc.add_paragraph(), p.replace("\n", " "))
        # figures and tables first mentioned in this paragraph, inserted after it
        for fid in [m.group(1) for m in FIG_RE.finditer(p)]:
            if fid in seen_f or fig_file(fid) is None:
                continue
            seen_f.append(fid); f = fig_file(fid); img = small_image(f, fid); n_fig += 1
            cap = caps.get(fid, f"**{fid}.**")
            doc.add_picture(str(img), width=Cm(16.5)); cp = doc.add_paragraph(); runs(cp, cap + (f" *({cap_note})*" if cap_note else "")); cp.runs[0].font.size = Pt(9)
            for r in cp.runs: r.font.size = Pt(9)
            md += [f"![{fid}](img/{fid}.jpg)", cap + (f" *({cap_note})*" if cap_note else "")]
        for tid in [m.group(0) for m in TAB_RE.finditer(p)]:
            if tid in seen_t or tid not in tman or not (TABS / f"{tid}.csv").exists():
                continue
            seen_t.append(tid); n_tab += 1; df, note = table_block(tid, a.lang)
            cap = f"**{tab_word} {tid}.** {tman[tid].get('caption', '')}" + (f" [{tman[tid].get('evidence_level', '')}]" if tman[tid].get("evidence_level") else "") + (f" *({note})*" if note else "")
            cp = doc.add_paragraph(); runs(cp, cap)
            for r in cp.runs: r.font.size = Pt(9)
            tb = doc.add_table(rows=1, cols=len(df.columns)); tb.style = "Table Grid"
            for j, c in enumerate(df.columns):
                cell = tb.rows[0].cells[j]; cell.text = str(c)
                for r in cell.paragraphs[0].runs: r.font.size = Pt(7); r.bold = True
            for row in df.itertuples(index=False):
                cells = tb.add_row().cells
                for j, v in enumerate(row):
                    cells[j].text = str(v)[:80]
                    for r in cells[j].paragraphs[0].runs: r.font.size = Pt(7)
            doc.add_paragraph()
            md += [cap, md_table(df)]
    stem = f"manuscript_{a.lang}_assembled"
    doc.save(OUT / f"{stem}.docx")
    head = (f"<!-- {stem}.md: assembled by workflows/paper/p101_assemble_doc.py from {src.name}; figures from publication/figures (downscaled copies in img/), "
            f"tables from publication/tables (first {MAX_ROWS} rows x {MAX_COLS} columns), captions from captions.md and tables/manifest.json. The filled manuscript is the source of truth. -->\n\n")
    (OUT / f"{stem}.md").write_text(head + "\n\n".join(md) + "\n", encoding="utf-8")
    size = (OUT / f"{stem}.docx").stat().st_size
    print(f"-> {OUT / stem}.docx ({size / 1e6:.1f} MB) and .md: {n_fig} figures, {n_tab} tables inserted; figures without a file: "
          f"{sorted(set(m.group(1) for m in FIG_RE.finditer(text)) - set(seen_f))}")


if __name__ == "__main__":
    main()
